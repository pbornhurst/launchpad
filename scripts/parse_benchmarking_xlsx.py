#!/usr/bin/env python3
"""
parse_benchmarking_xlsx.py — Parse a DoorDash competitive benchmarking xlsx
export into structured JSON, suitable for ingestion by the /benchmark-doc agent.

Usage:
  python3 parse_benchmarking_xlsx.py "/path/to/Combined Version.xlsx"
  python3 parse_benchmarking_xlsx.py "/path/to/Combined Version.xlsx" --include-raw

Writes one JSON object to stdout. Top-level keys:
  identity        — store address, cuisine, starting point name
  report_month    — inferred latest month present in the workbook (YYYY-MM)
  headline        — sales totals, MoM/QoQ/YoY growth, rank in category & area
  visibility      — impressions / CTR / CVR comparison and monthly trends by source
  search          — top 10 store search terms + area top 25 + area trending
  customers       — cohort counts, cohort mix vs peers, top 50 customers
  competition     — top competitor menu items, active area campaigns
  operations      — operational quality (prep, rating, photos, promo/ads spend)
  geography       — top zips with reach/penetration
  raw_sheets      — (only with --include-raw) full row dump per sheet
"""

import json
import sys
from datetime import datetime, date
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.stderr.write("openpyxl is required: pip3 install openpyxl\n")
    sys.exit(2)


# ---- helpers ----------------------------------------------------------------

def _cell(v):
    if isinstance(v, (datetime, date)):
        if isinstance(v, datetime) and v.time() == datetime.min.time():
            return v.date().isoformat()
        return v.isoformat()
    return v


def _sheet_rows(ws):
    return [[_cell(c) for c in row] for row in ws.iter_rows(values_only=True)]


def _records(ws, header_row=0):
    rows = _sheet_rows(ws)
    if not rows or len(rows) <= header_row:
        return []
    headers = [str(h) if h is not None else f"col_{i}" for i, h in enumerate(rows[header_row])]
    out = []
    for r in rows[header_row + 1:]:
        if all(c is None for c in r):
            continue
        record = {}
        for i, val in enumerate(r):
            key = headers[i] if i < len(headers) else f"col_{i}"
            record[key] = val
        out.append(record)
    return out


def _pivot_records(ws):
    """For sheets whose row 1 is a title and row 2 is the real header."""
    return _records(ws, header_row=1)


def _kv(ws):
    """For 2-row sheets that are label/value pairs."""
    rows = _sheet_rows(ws)
    if len(rows) < 2:
        return {}
    return {str(h): v for h, v in zip(rows[0], rows[1])}


def _find_sheet(wb, *candidates):
    """Resolve tab names with tolerance for trailing spaces and truncation."""
    names = wb.sheetnames
    for cand in candidates:
        if cand in names:
            return wb[cand]
    # Fuzzy: prefix match after stripping
    for cand in candidates:
        c = cand.strip().lower()
        for n in names:
            if n.strip().lower().startswith(c[:28]):
                return wb[n]
    return None


def _months_in_sheet(ws):
    """Pull all month values out of column A (skipping headers)."""
    months = []
    for row in ws.iter_rows(min_row=1, values_only=True):
        if not row:
            continue
        v = row[0]
        if isinstance(v, (datetime, date)):
            months.append(v.date() if isinstance(v, datetime) else v)
    return months


# ---- main parse -------------------------------------------------------------

def parse(path, include_raw=False):
    wb = openpyxl.load_workbook(path, data_only=True)

    out = {
        "source_file": str(path),
        "parsed_at": datetime.now().isoformat(timespec="seconds"),
    }

    # Identity
    out["identity"] = {
        "store_address": _kv(_find_sheet(wb, "Store Address")).get("Store Address"),
        "cuisine": _kv(_find_sheet(wb, "Cuisine")).get("Cuisine"),
        "starting_point": _kv(_find_sheet(wb, "Starting Point Name")).get("Starting Point Name"),
    }

    # Report month inference — latest month across the time-series sheets
    candidate_sheets = ["% Impressions by Source", "CTR by Source", "CVR by Source", "# Cxs by Cohort"]
    all_months = []
    for s in candidate_sheets:
        ws = _find_sheet(wb, s)
        if ws is not None:
            all_months.extend(_months_in_sheet(ws))
    if all_months:
        latest = max(all_months)
        out["report_month"] = latest.strftime("%Y-%m")
        out["report_month_label"] = latest.strftime("%B %Y")
    else:
        out["report_month"] = None
        out["report_month_label"] = None

    # Headline KPIs
    headline = {}
    sc = _find_sheet(wb, "Sales Comparison")
    if sc is not None:
        headline["sales_comparison"] = _records(sc)
    headline["sales_of_report_month"] = _kv(_find_sheet(wb, "Sales of Report Month")).get("Report Month")
    headline["sales_growth_mom"] = _kv(_find_sheet(wb, "Sales Growth MM")).get("Mom")
    headline["category_total_sales"] = _kv(_find_sheet(wb, "Category & Delivery Area Sal", "Category & Delivery Area Sales")).get("Sum of Category Sales")
    headline["delivery_area_total_sales"] = _kv(_find_sheet(wb, "Delivery Area Sales")).get("Sales in Sp")

    rk_cat = _find_sheet(wb, "Sales Rank in Your CategoryA", "Sales Rank in Your Category")
    if rk_cat is not None:
        headline["category_rank"] = _kv(rk_cat)
    rk_area = _find_sheet(wb, "Sales Rank in Your Delivery ", "Sales Rank in Your Delivery Area")
    if rk_area is not None:
        headline["delivery_area_rank"] = _kv(rk_area)

    cat_perf = _find_sheet(wb, "Category & Delivery Area Per", "Category & Delivery Area Performance")
    if cat_perf is not None:
        recs = _records(cat_perf)
        headline["category_performance"] = recs[0] if recs else None
    area_perf = _find_sheet(wb, "Delivery Area Performance Pe", "Delivery Area Performance")
    if area_perf is not None:
        recs = _records(area_perf)
        headline["delivery_area_performance"] = recs[0] if recs else None
    out["headline"] = headline

    # Visibility funnel
    visibility = {}
    vis = _find_sheet(wb, "Visibility Comparison")
    if vis is not None:
        visibility["comparison"] = _records(vis)
    imp = _find_sheet(wb, "Impression by Source")
    if imp is not None:
        visibility["impression_share_by_source"] = _records(imp)
    ctr = _find_sheet(wb, "CTR  by Source", "CTR by Source vs Peers")
    if ctr is not None:
        visibility["ctr_by_source"] = _records(ctr)
    cvr = _find_sheet(wb, "CVR by Source-1", "CVR by Source vs Peers")
    if cvr is not None:
        visibility["cvr_by_source"] = _records(cvr)
    # Time-series (pivot style)
    pct_imp_ts = _find_sheet(wb, "% Impressions by Source")
    if pct_imp_ts is not None:
        visibility["pct_impressions_monthly"] = _pivot_records(pct_imp_ts)
    ctr_ts = _find_sheet(wb, "CTR by Source")
    if ctr_ts is not None:
        visibility["ctr_monthly"] = _pivot_records(ctr_ts)
    cvr_ts = _find_sheet(wb, "CVR by Source")
    if cvr_ts is not None:
        visibility["cvr_monthly"] = _pivot_records(cvr_ts)
    agg = _find_sheet(wb, "Aggregate Impressions by Sou", "Aggregate Impressions by Source")
    if agg is not None:
        visibility["aggregate_impressions"] = _records(agg)
    out["visibility"] = visibility

    # Search
    search = {}
    top10 = _find_sheet(wb, "Your Top 10 Search Terms")
    if top10 is not None:
        search["your_top_terms"] = _records(top10)
    top25 = _find_sheet(wb, "Most Popular 25 Searches in ", "Most Popular 25 Searches")
    if top25 is not None:
        search["area_top_searches"] = _records(top25)
    trend = _find_sheet(wb, " Trending (Highest MM Growth", "Trending (Highest MM Growth)")
    if trend is not None:
        search["area_trending"] = _records(trend)
    out["search"] = search

    # Customers
    customers = {}
    cnt = _find_sheet(wb, "# Cxs by Cohort")
    if cnt is not None:
        customers["counts_monthly"] = _pivot_records(cnt)
    mom = _find_sheet(wb, "# Cxs MM by Cohort")
    if mom is not None:
        customers["counts_mom"] = _pivot_records(mom)
    ctype = _find_sheet(wb, "% Orders by Customer Type")
    if ctype is not None:
        customers["pct_orders_by_type"] = _records(ctype)
    cmix = _find_sheet(wb, "% Orders by Cohort Compariso", "% Orders by Cohort Comparison")
    if cmix is not None:
        customers["cohort_mix_vs_peers"] = _records(cmix)
    top50 = _find_sheet(wb, "Your Top 50 Customers")
    if top50 is not None:
        customers["top_customers"] = _records(top50)
    out["customers"] = customers

    # Competition
    competition = {}
    peer_items = _find_sheet(wb, "Top items offered by competi", "Top items offered by competitors")
    if peer_items is not None:
        competition["peer_top_items"] = _records(peer_items)
    camps = _find_sheet(wb, "Campaigns being run in your ", "Campaigns being run in your area")
    if camps is not None:
        competition["area_campaigns"] = _records(camps)
    out["competition"] = competition

    # Operations
    operations = {}
    opq = _find_sheet(wb, "Operational Quality")
    if opq is not None:
        operations["quality"] = _records(opq)
    out["operations"] = operations

    # Geography
    geography = {}
    nm = _find_sheet(wb, "New Map")
    if nm is not None:
        geography["postal_deliveries"] = _records(nm)
    zips = _find_sheet(wb, "Orders by Zip ", "Orders by Zip")
    if zips is not None:
        geography["orders_by_zip"] = _records(zips)
    out["geography"] = geography

    # Optional raw fallback
    if include_raw:
        raw = {}
        for s in wb.sheetnames:
            raw[s] = _sheet_rows(wb[s])
        out["raw_sheets"] = raw

    return out


def main():
    args = [a for a in sys.argv[1:] if a not in ("--include-raw",)]
    include_raw = "--include-raw" in sys.argv[1:]
    if not args:
        sys.stderr.write("usage: parse_benchmarking_xlsx.py <path-to-xlsx> [--include-raw]\n")
        sys.exit(2)
    path = Path(args[0]).expanduser()
    if not path.exists():
        sys.stderr.write(f"file not found: {path}\n")
        sys.exit(1)
    data = parse(path, include_raw=include_raw)
    sys.stdout.write(json.dumps(data, default=str, indent=None))
    sys.stdout.write("\n")


if __name__ == "__main__":
    main()
